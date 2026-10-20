"""Build the Employment Impacts and Productivity Impacts data files from the "Main Data for Sharing" workbook.

Writes, for each of the two tabs:
    data/employment-impacts-data.js     data/employment-impacts.csv
    data/productivity-impacts-data.js   data/productivity-impacts.csv

Usage (from anywhere; requires openpyxl):
    python3 scripts/build_impact_data.py [path/to/workbook.xlsx]

The workbook defaults to WORKBOOK below. Each impact sheet ("AI Emp Past 3y", "AI Emp Next 3y", "AI Prod Past 3y",
"AI Prod Next 3y") holds one table per survey wave; WAVES says where each table starts and gives the fieldwork dates
and notes, which are not in the tables themselves. To add a new wave, append it to WAVES (latest last) and update the
table positions. The script checks every table against the workbook's own formulas before writing anything.
"""
import json
import sys
from pathlib import Path

import openpyxl

REPO = Path(__file__).resolve().parent.parent
WORKBOOK = REPO / "data" / "Firm Data on AI - Main Data for Sharing - October 2026.xlsx"

# (id, label, survey, column header in the workbook), in display order.
COUNTRIES = [
    ("all_firms", "All Firms", "Average across all firms in the surveyed countries", "All Firms"),
    ("us", "US", "US Survey of Business Uncertainty", "US Survey of Business Uncertainty"),
    ("canada", "Canada", "Canada Business Leaders' Pulse (BLP)", "Canada Business Leaders' Pulse (BLP)"),
    ("uk", "UK", "UK Decision Maker Panel", "UK Decision Maker Panel"),
    ("germany", "Germany", "Germany Bundesbank Online Panel - Firms", "Germany Bundesbank Online Panel - Firms"),
    ("australia", "Australia", "Australia Business Outlook Scenarios Survey", "Australia Business Outlook Scenarios Survey"),
    ("mexico", "Mexico", "Mexico Monthly Survey on Regional Economic Activity (EMAER)", "Mexico Monthly Survey on Regional Economic Activity (EMAER)"),
    ("south_africa", "South Africa", "South Africa Quarterly inflation expectations survey", "South Africa Quarterly inflation expectations survey"),
]
HEADER_TO_ID = {header: cid for cid, _, _, header in COUNTRIES}

# Survey waves in chronological order. "headerRow" is the row holding the country names of that wave's table, the
# same in every impact sheet. A footnote puts an asterisk on that country in the charts and table.
WAVES = [
    {
        "id": "2025_nov_jan",
        "label": "Nov 2025 – Jan 2026",
        "shortLabel": "Nov 25–Jan 26",
        "headerRow": 23,
        "fieldwork": {
            "all_firms": "Nov 2025 – Jan 2026",
            "us": "Nov 2025",
            "uk": "Nov 2025 – Jan 2026",
            "germany": "Jan 2026",
            "australia": "Dec 2025",
        },
        "notes": (
            "The data from the US Survey of Business Uncertainty was collected in November 2025. The data from the UK "
            "Decision Maker Panel was collected over November 2025 – January 2026. The data from the German Bundesbank "
            "Online Panel – Firms (BOP-F) were collected in January 2026. The data from the Australian Business Outlook "
            "Scenarios Survey was collected in December 2025. The results from the SBU, DMP, and BOP-F are "
            "employment-weighted; the results from the BOSS are unweighted."
        ),
    },
    {
        "id": "2026_may_jul",
        "label": "May – Jul 2026",
        "shortLabel": "May–Jul 26",
        "headerRow": 3,
        "fieldwork": {
            "all_firms": "May – Jul 2026",
            "us": "Jun 2026",
            "uk": "May – Jul 2026",
            "germany": "Jan 2026",
            "australia": "Jun 2026",
            "canada": "May – Jun 2026",
            "mexico": "Jun 2026",
            "south_africa": "May 2026",
        },
        "footnotes": {"germany": "Germany data were collected in January 2026; no newer data on AI impacts are available."},
        "notes": (
            "The data from the US Survey of Business Uncertainty was collected in June 2026. The data from the UK "
            "Decision Maker Panel was collected over May – July 2026. The data from the German Bundesbank Online Panel – "
            "Firms (BOP-F) were collected in January 2026, as no newer data on AI impacts are available. The data from the "
            "Australian Business Outlook Scenarios Survey was collected in June 2026. The data from the Canadian Business "
            "Leaders' Pulse was collected over May – June 2026. The data from the Mexican Monthly Survey on Regional "
            "Economic Activity was collected in June 2026. The data from the South African Quarterly inflation "
            "expectations survey was collected in May 2026. The results from the SBU, DMP, BOP-F, and EMAER are "
            "employment-weighted; the results from the BOSS, BLP, and South African Quarterly inflation expectations "
            "survey are equal-weighted."
        ),
    },
]

# Values assigned to the five answer options (large negative ... large positive) for the cumulative impact.
WEIGHTS = [-7.5, -2.5, 0, 2.5, 7.5]

DATASETS = {
    "employmentImpactsData": {
        "id": "employment_impacts",
        "tabLabel": "Employment Impacts",
        "downloadBase": "employment-impacts-ai",
        "file": "employment-impacts",
        "measure": "employment",
        "subject": "the number of employees",
        "sheets": ("AI Emp Past 3y", "AI Emp Next 3y"),
    },
    "productivityImpactsData": {
        "id": "productivity_impacts",
        "tabLabel": "Productivity Impacts",
        "downloadBase": "productivity-impacts-ai",
        "file": "productivity-impacts",
        "measure": "productivity",
        "subject": "the volume of sales per employee",
        "sheets": ("AI Prod Past 3y", "AI Prod Next 3y"),
    },
}


def question(subject, horizon):
    past = f"How has the adoption of artificial intelligence technologies affected {subject} in your business over the past three years?"
    if horizon == "past":
        return f'The data are based on responses to the question: "{past}"'
    return (
        f'The data are based on responses to the question: "{past} And how do you expect this to affect '
        f'{subject.replace("the ", "your ", 1)} over the next 3 years?"'
    )


def build_periods(cfg):
    measure = cfg["measure"]
    return [
        {
            "id": "past_3_years",
            "label": "Impacts over last 3 years",
            "sheet": cfg["sheets"][0],
            "title": f"Impact of AI on {measure} over the past 3 years (% of respondents, cumulative impact)",
            "horizon": "over the past 3 years",
            "averageLabel": "Cumulative impact over past 3 years (%)",
            "questionNote": question(cfg["subject"], "past"),
        },
        {
            "id": "next_3_years",
            "label": "Expected impacts over next 3 years",
            "sheet": cfg["sheets"][1],
            "title": f"Expected impact of AI on {measure} over the next 3 years (% of respondents, cumulative impact)",
            "horizon": "over the next 3 years",
            "averageLabel": "Cumulative impact over next 3 years (%)",
            "questionNote": question(cfg["subject"], "next"),
        },
    ]


def check(condition, message):
    if not condition:
        sys.exit(f"Workbook check failed: {message}")


def number(ws, row, col, where):
    value = ws.cell(row, col).value
    check(
        isinstance(value, (int, float)),
        f"{where}: cell {ws.cell(row, col).coordinate} is {value!r}, not a number. If it is a formula, open and save "
        "the workbook in Excel so the calculated values are stored.",
    )
    return value


def read_categories(workbook, sheet, wave):
    """The five answer options labelling a wave's table, with stray spaces collapsed."""
    ws = workbook[sheet]
    h = wave["headerRow"]
    return [" ".join(str(ws.cell(h + 1 + i, 1).value).split()) for i in range(5)]


def read_wave(workbook, sheet, wave):
    ws = workbook[sheet]
    h = wave["headerRow"]
    where = f'{sheet}, {wave["label"]} table (row {h})'
    check(ws.cell(h + 6, 1).value == "Number of Survey Responses", f"{where}: responses not found at row {h + 6}")
    check(str(ws.cell(h + 9, 1).value).startswith("Cumulative impact"), f"{where}: cumulative impact not found at row {h + 9}")

    data = {}
    for col in range(2, ws.max_column + 1):
        header = ws.cell(h, col).value
        if not header:
            continue
        check(header in HEADER_TO_ID, f"{where}: unknown column {header!r}")
        cid = HEADER_TO_ID[header]
        check(cid in wave["fieldwork"], f"{where}: no fieldwork date for {cid} in WAVES")
        values = [float(number(ws, h + 1 + i, col, where)) for i in range(5)]
        responses = int(round(number(ws, h + 6, col, where)))
        cumulative = float(number(ws, h + 9, col, where))
        implied = sum(w * v for w, v in zip(WEIGHTS, values)) / 100
        check(abs(implied - cumulative) < 1e-3, f"{where}: {cid} cumulative impact {cumulative} != {implied} implied by the shares")
        data[cid] = {"fieldwork": wave["fieldwork"][cid], "responses": responses, "values": values, "average": cumulative}
        if cid in wave.get("footnotes", {}):
            data[cid]["footnote"] = wave["footnotes"][cid]

    # "All Firms" should be the response-weighted average of the surveys in the wave.
    check("all_firms" in data, f"{where}: no All Firms column")
    countries = [cid for cid in data if cid != "all_firms"]
    total = sum(data[c]["responses"] for c in countries)
    check(total == data["all_firms"]["responses"], f"{where}: All Firms responses {data['all_firms']['responses']} != {total}")
    for i in range(5):
        weighted = sum(data[c]["values"][i] * data[c]["responses"] for c in countries) / total
        check(abs(weighted - data["all_firms"]["values"][i]) < 1e-3, f"{where}: All Firms row {i + 1} is not the weighted average")
    weighted = sum(data[c]["average"] * data[c]["responses"] for c in countries) / total
    check(abs(weighted - data["all_firms"]["average"]) < 1e-3, f"{where}: All Firms cumulative impact is not the weighted average")

    return {
        "id": wave["id"],
        "label": wave["label"],
        "shortLabel": wave["shortLabel"],
        "data": {cid: data[cid] for cid, *_ in COUNTRIES if cid in data},
        "notes": wave["notes"],
    }


# ---- JavaScript data module ----

def js_list(value, indent):
    return json.dumps(value, ensure_ascii=False, indent=2).replace("\n", "\n" + " " * indent)


def entry_js(entry):
    return (
        f'{{ fieldwork: {json.dumps(entry["fieldwork"], ensure_ascii=False)}, responses: {entry["responses"]}, '
        f'values: [{", ".join(repr(v) for v in entry["values"])}], average: {entry["average"]!r}'
        + (f', footnote: {json.dumps(entry["footnote"], ensure_ascii=False)}' if "footnote" in entry else "")
        + " }"
    )


def period_js(period):
    lines = [
        "    {",
        f'      id: "{period["id"]}",',
        f'      label: "{period["label"]}",',
        f'      sourceSheet: "{period["sheet"]}",',
        f"      title: {json.dumps(period['title'])},",
        f'      horizon: "{period["horizon"]}",',
        f'      averageLabel: "{period["averageLabel"]}",',
        "      questionNote:",
        f"        {json.dumps(period['questionNote'], ensure_ascii=False)},",
        "      waves: [",
    ]
    wave_blocks = []
    for wave in period["waves"]:
        block = [
            "        {",
            f'          id: "{wave["id"]}",',
            f'          label: "{wave["label"]}",',
            f'          shortLabel: "{wave["shortLabel"]}",',
            "          data: {",
            ",\n".join(f"            {cid}: {entry_js(entry)}" for cid, entry in wave["data"].items()),
            "          },",
            "          notes:",
            f"            {json.dumps(wave['notes'], ensure_ascii=False)}",
            "        }",
        ]
        wave_blocks.append("\n".join(block))
    lines.append(",\n".join(wave_blocks))
    lines += ["      ]", "    }"]
    return "\n".join(lines)


def dataset_js(var, cfg, periods, categories):
    measure = cfg["measure"]
    periods_js = ",\n".join(period_js(period) for period in periods)
    countries = ",\n".join(
        f'    {{ id: "{cid}", label: "{label}", source: {json.dumps(source)} }}' for cid, label, source, _ in COUNTRIES
    )
    return f"""const {var} = {{
  id: "{cfg["id"]}",
  title: "Impact of AI on {measure}",
  tabLabel: "{cfg["tabLabel"]}",
  rowLabel: "Impact Category",
  downloadBase: "{cfg["downloadBase"]}",
  measureLabel: "{measure}",
  defaultPeriod: "past_3_years",
  categories: {js_list(categories, 2)},
  // Display order. "All Firms" in each wave is the response-weighted average of the countries in that wave.
  countries: [
{countries}
  ],
  // One entry per impact horizon, each with its survey waves in chronological order, earliest first. Charts show
  // the latest wave; the change chart compares all waves. Each entry's "average" is its cumulative impact (%).
  periods: [
{periods_js}
  ],
  methodNote:
    "The results for all firms are the average of the surveys in each wave, weighted by the respective number of responses. To calculate the cumulative impact, values are assigned to each of the options: large negative/large positive impacts are treated as ±7.5%; small negative/small positive impacts are treated as ±2.5%."
}};
"""


# ---- CSV, matching the dashboard's "Download CSV" output (waveDatasetToCsv in app.js) ----

def js_number(value):
    """Format a number the way JavaScript's Number#toString does for these values."""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return repr(value)


def csv_cell(cell):
    if isinstance(cell, (int, float)):
        return js_number(cell)
    escaped = str(cell).replace('"', '""')
    return f'"{escaped}"' if any(ch in escaped for ch in '",\n') else escaped


def panel_countries(waves):
    """Countries with new survey data in every wave (getPanelCountries in app.js)."""
    ids = []
    for cid, *_ in COUNTRIES:
        if cid == "all_firms" or not all(cid in wave["data"] for wave in waves):
            continue
        if len({wave["data"][cid]["fieldwork"] for wave in waves}) == len(waves):
            ids.append(cid)
    return ids


def weighted_entry(wave, ids):
    """Response-weighted average over a subset of countries (computeWeightedEntry in app.js)."""
    entries = [wave["data"][cid] for cid in ids]
    responses = sum(entry["responses"] for entry in entries)

    def weighted(pick):
        total = 0
        for entry in entries:
            total += pick(entry) * entry["responses"]
        return total / responses

    return {
        "fieldwork": wave["label"],
        "responses": responses,
        "values": [weighted(lambda entry, i=i: entry["values"][i]) for i in range(5)],
        "average": weighted(lambda entry: entry["average"]),
    }


def format_list(items):
    return items[0] if len(items) == 1 else f'{", ".join(items[:-1])} and {items[-1]}'


def dataset_csv(periods, categories):
    labels = {cid: (label, source) for cid, label, source, _ in COUNTRIES}
    rows = [["Horizon", "Wave", "Country", "Survey", "Fieldwork", "Number of Survey Responses", "Impact Category", "Value"]]
    for period in periods:
        names = [f"{name} (%)" for name in categories] + [period["averageLabel"]]
        panel = panel_countries(period["waves"])

        def push(wave, label, source, entry):
            for name, value in zip(names, entry["values"] + [entry["average"]]):
                rows.append([period["label"], wave["label"], label, source, entry["fieldwork"], entry["responses"], name, value])

        for wave in period["waves"]:
            for cid, entry in wave["data"].items():
                push(wave, *labels[cid], entry)
            # Like-for-like all-firms figure for the change chart, when the wave covers other countries too.
            if len(wave["data"]) - 1 != len(panel):
                label = f'All Firms ({format_list([labels[cid][0] for cid in panel])} only)'
                push(wave, label, "Response-weighted average of the countries surveyed in all waves", weighted_entry(wave, panel))
    return "\n".join(",".join(csv_cell(cell) for cell in row) for row in rows) + "\n"


def main():
    workbook_path = Path(sys.argv[1]) if len(sys.argv) > 1 else WORKBOOK
    workbook = openpyxl.load_workbook(workbook_path, data_only=True)

    outputs = {}
    for var, cfg in DATASETS.items():
        periods = build_periods(cfg)
        for period in periods:
            period["waves"] = [read_wave(workbook, period["sheet"], wave) for wave in WAVES]

        # One set of answer options is shared by every horizon and wave, so the tables must agree on it.
        labels = {
            f'{period["sheet"]}, {wave["label"]}': read_categories(workbook, period["sheet"], wave)
            for period in periods
            for wave in WAVES
        }
        categories = next(iter(labels.values()))
        for where, found in labels.items():
            check(found == categories, f"{where}: answer options {found} differ from {categories}")

        outputs[REPO / "data" / f'{cfg["file"]}-data.js'] = dataset_js(var, cfg, periods, categories)
        outputs[REPO / "data" / f'{cfg["file"]}.csv'] = dataset_csv(periods, categories)

    # Only write once every sheet has passed its checks.
    for path, text in outputs.items():
        path.write_text(text, encoding="utf-8")
        print(f"wrote {path.relative_to(REPO)}")


if __name__ == "__main__":
    main()
